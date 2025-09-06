import os
import re
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# 文件夹路径，存放多个xlsx文件（含黄色标记的）
folder_path = "./reclassify_result_by_xinyue"
# 需要删除行的目标Excel文件
target_file = "./classify_results.xlsx"

# 定义黄色填充样式判断
YELLOW_RGBS = {"FFFF00", "FFFFFF00"}

# 用来存储所有黄色标记的 (第一列内容去括号, 第二列内容) 元组
yellow_tuples = set()

# 匹配中文括号（）
bracket_pattern = re.compile(r'（.*?）')

# 遍历文件夹里的所有xlsx文件
for filename in os.listdir(folder_path):
    if filename.endswith(".xlsx"):
        filepath = os.path.join(folder_path, filename)
        wb = load_workbook(filepath)
        sheet = wb.active

        for row in sheet.iter_rows():
            for cell in row:
                fill_color = getattr(cell.fill.start_color, 'rgb', None)
                if fill_color in YELLOW_RGBS:
                    # 取第一列和第二列
                    col1 = row[0].value
                    col2 = row[1].value
                    if col1 is None or col2 is None:
                        break
                    # 去除中文括号及其内容
                    col1_clean = re.sub(bracket_pattern, "", str(col1))
                    # 保存为元组
                    yellow_tuples.add((col1_clean, str(col2)))
                    break

print("黄色标记提取完成，共找到：", len(yellow_tuples))

# 打开目标文件
target_wb = load_workbook(target_file)
target_ws = target_wb.active

# 需要删除的行索引（注意不能边遍历边删）
rows_to_delete = []
for row in target_ws.iter_rows(min_row=2):  # 跳过表头（如果有表头）
    col1 = row[0].value
    col2 = row[1].value
    if col1 is None or col2 is None:
        continue
    col1_clean = re.sub(bracket_pattern, "", str(col1))
    if (col1_clean, str(col2)) in yellow_tuples:
        rows_to_delete.append(row[0].row)

# 倒序删除行，避免行号错位
for row_idx in sorted(rows_to_delete, reverse=True):
    target_ws.delete_rows(row_idx)

print("删除完成，共删除：", len(rows_to_delete), "行")

# 保存
target_wb.save("./target_cleaned.xlsx")
print("保存到 target_cleaned.xlsx")
